"""legacy.py -- the 20 legacy (2019) pick-glass images with partial hand counts.

What lives here
---------------
* ``load_gt()``          : parse "Warp-Weft count.xlsx" (top-level copy) into tidy records.
* ``load_baseline2019()``: the 2019 "predicted weft by RGB irregularity" column that only
                           exists in the copy under "FR-final attempt/" (a 2019 baseline).
* ``detect_window(img)`` : simple, documented frame detector for the legacy crops.
* ``rectify_legacy(name)``: square crop of the detected 1-inch window (S x S, S = mean side).
* ``duplicate_check()``  : perceptual-hash comparison of legacy crops vs. the 150-image set.

Legacy images are *already hand-cropped* close to the 1-inch window (about 1650-2010 px
square) but a sliver of the black pick-glass frame with its bright metal tick marks is
visible on some sides, and on other sides the crop cuts slightly inside the window.

Window detector (``detect_window``)
-----------------------------------
1. V = max(B,G,R), Gaussian-smoothed (sigma = 2 px).  Reference brightness ``ref`` = median V
   of the central 50 % of the image (always fabric).
2. A pixel is *frame-like* if V < 0.30*ref (black plastic frame and its shadow; the thin
   bright tick marks only lower the dark fraction of the frame band, they do not break it).
   Legacy fabrics are all mid/bright, so this is unambiguous here; it would NOT work for
   very dark fabrics (documented limitation).
3. For every side, the side is split into 24 strips (corners excluded). Walking inwards from
   the image border (at most 22 % of the side), the edge is the first offset after which the
   strip's frame-like fraction stays below 0.35 for 12 consecutive pixel rows/cols.  If the
   border row is already fabric, the edge is 0 (the crop cut inside the window; the true
   window is then slightly larger than the crop -> counts on that axis are biased low).
4. A Theil-Sen line is fitted to the 24 edge points of each side (robust to tick marks;
   if |slope| > 0.035 the side falls back to a horizontal/vertical line at the median offset),
   the four lines are intersected -> TL, TR, BR, BL, and the quad is warped to a square
   of side round(mean side length), i.e. the same convention as common/rectify.py.
"""
from __future__ import annotations

import datetime as _dt
import os
import re

import cv2
import numpy as np
from scipy import stats

ROOT = (os.environ.get("PICKGLASS_LEGACY_ROOT") or os.path.join(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."), "data", "legacy2019"))
LEGACY_DIR = os.path.join(ROOT, "FR-final attempt", "images")
GT_XLSX = os.path.join(ROOT, "Warp-Weft count.xlsx")
GT_XLSX_2019PRED = os.path.join(ROOT, "FR-final attempt", "Warp-Weft count.xlsx")
NAMES = [f"a{i}" for i in range(1, 11)] + [f"r{i}" for i in range(1, 11)]


def legacy_path(name: str) -> str:
    p = os.path.join(LEGACY_DIR, f"{name}.jpg")
    if not os.path.exists(p):  # "r6 .jpg" has a space before the extension
        p = os.path.join(LEGACY_DIR, f"{name} .jpg")
    return p


def load_legacy(name: str) -> np.ndarray:
    img = cv2.imread(legacy_path(name), cv2.IMREAD_COLOR)
    if img is None:
        raise FileNotFoundError(legacy_path(name))
    return img


# ----------------------------------------------------------------------------- GT parsing
def _parse_count(v):
    """Returns (value, flag). flag in {'', 'formula', 'range'}; value None if missing."""
    if v is None:
        return None, ""
    if isinstance(v, (int, float)):
        return float(v), ""
    s = str(v).strip()
    if "=" in s:  # e.g. "26 + 25*2 = 76" -> 76
        return float(s.split("=")[-1].strip()), "formula"
    m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)", s)
    if m:  # "110-120" -> midpoint, flagged
        return 0.5 * (float(m.group(1)) + float(m.group(2))), "range"
    try:
        return float(s), ""
    except ValueError:
        return None, "unparsed:" + s


def _parse_weave(v):
    """Excel turned '1/1' into 2019-01-01 and '2/1' into 2019-02-01 (US M/D)."""
    if v is None:
        return None
    if isinstance(v, (_dt.datetime, _dt.date)):
        return f"{v.month}/{v.day}"
    return str(v).strip()


WEAVE_MAP = {"1/1": "plain", "2/1": "twill_2_1", "1/2": "twill_2_1", "3/1": "twill_3_1",
             "2/2": "twill_2_2"}


def load_gt(path: str = GT_XLSX):
    """List of dicts: name, unclear, weft, warp, warp_flag, weft_flag, interlacing, weave."""
    import openpyxl

    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    col = {h: k for k, h in enumerate(hdr)}
    out = []
    for r in rows[1:]:
        raw = r[col["img file name"]]
        if raw is None or not str(raw).strip() or str(raw).startswith("* image"):
            continue
        raw = str(raw).strip()
        name = raw.rstrip("*").strip()
        weft, wf = _parse_count(r[col["weft count"]])
        warp, pf = _parse_count(r[col["warp count"]])
        inter = _parse_weave(r[col["interlacing pattern"]])
        out.append(dict(name=name, unclear=raw.endswith("*"), weft=weft, warp=warp,
                        weft_flag=wf, warp_flag=pf, interlacing=inter,
                        weave=WEAVE_MAP.get(inter) if inter else None,
                        has_counts=(weft is not None and warp is not None)))
    return out


def load_baseline2019(path: str = GT_XLSX_2019PRED):
    """{'a1': 69.0, ...}: 2019 pipeline's 'predicted weft by RGB irregularity' (weft only)."""
    import openpyxl

    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    k = [i for i, h in enumerate(hdr) if h.startswith("predicted weft")]
    if not k:
        return {}
    out = {}
    for r in rows[1:]:
        if r[0] is None or str(r[0]).startswith("* image"):
            continue
        name = str(r[0]).strip().rstrip("*").strip()
        if r[k[0]] is not None:
            out[name] = float(r[k[0]])
    return out


# ------------------------------------------------------------------------ window detection
def _edge_offsets(frame: np.ndarray, side: str, n_strips=24, max_frac=0.22,
                  clear_frac=0.35, run=12):
    """Edge offset (px from the border) per strip for one side; frame = bool frame-like map."""
    H, W = frame.shape
    if side in ("left", "right"):
        f = frame.T  # make the side the "top" of a (W x H) array
        L, N = W, H
    else:
        f = frame
        L, N = H, W
    if side in ("bottom", "right"):
        f = f[::-1]
    lim = int(max_frac * L)
    lo, hi = int(0.08 * N), int(0.92 * N)
    edges = np.linspace(lo, hi, n_strips + 1).astype(int)
    pos, offs = [], []
    for a, b in zip(edges[:-1], edges[1:]):
        frac = f[:lim, a:b].mean(axis=1)
        frac = np.convolve(frac, np.ones(3) / 3, mode="same")
        clear = frac < clear_frac
        off = lim  # fallback: no clear run found
        # first index i such that clear[i:i+run] all True
        cs = np.concatenate([[0], np.cumsum(clear)])
        for i in range(0, lim - run):
            if cs[i + run] - cs[i] == run:
                off = i
                break
        pos.append(0.5 * (a + b))
        offs.append(off)
    return np.array(pos), np.array(offs, dtype=float), L


def detect_window(img: np.ndarray):
    """Return (corners TL,TR,BR,BL float (4,2), info dict) for a legacy hand-cropped image."""
    H, W = img.shape[:2]
    v = cv2.GaussianBlur(img.max(axis=2).astype(np.float32), (0, 0), 2)
    c = v[int(0.25 * H):int(0.75 * H), int(0.25 * W):int(0.75 * W)]
    ref = float(np.median(c))
    frame = v < 0.30 * ref
    lines = {}
    info = {"ref_brightness": ref}
    for side in ("top", "bottom", "left", "right"):
        pos, off, L = _edge_offsets(frame, side)
        ok = off < 0.21 * L
        if ok.sum() < 6:  # not enough evidence -> assume border
            slope, icpt = 0.0, 0.0
        else:
            slope, icpt, _, _ = stats.theilslopes(off[ok], pos[ok])
            if abs(slope) > 0.035:  # hand crops are near axis-aligned (<2 deg): fall back
                slope, icpt = 0.0, float(np.median(off[ok]))
        info[f"{side}_median_offset"] = float(np.median(off))
        lines[side] = (slope, icpt)

    # convert to image-coordinate lines: top: y = a x + b ; left: x = a y + b
    def top(x):
        a, b = lines["top"]; return a * x + b

    def bottom(x):
        a, b = lines["bottom"]; return (H - 1) - (a * x + b)

    def left(y):
        a, b = lines["left"]; return a * y + b

    def right(y):
        a, b = lines["right"]; return (W - 1) - (a * y + b)

    def inter(hline, vline):
        # solve y = hline(x), x = vline(y) by fixed-point iteration (lines are near-axis)
        x, y = W / 2, H / 2
        for _ in range(20):
            y = hline(x)
            x = vline(y)
        return [x, y]

    corners = np.array([inter(top, left), inter(top, right), inter(bottom, right),
                        inter(bottom, left)], dtype=np.float64)
    corners[:, 0] = np.clip(corners[:, 0], 0, W - 1)
    corners[:, 1] = np.clip(corners[:, 1], 0, H - 1)
    return corners, info


def rectify_quad(img: np.ndarray, corners, out_size=None) -> np.ndarray:
    """Same convention as common/rectify.py: S = round(mean side), bilinear warp."""
    c = np.asarray(corners, dtype=np.float32)
    S = int(out_size) if out_size else int(round(float(np.linalg.norm(c - np.roll(c, -1, 0), axis=1).mean())))
    dst = np.array([[0, 0], [S, 0], [S, S], [0, S]], dtype=np.float32)
    Hm = cv2.getPerspectiveTransform(c, dst)
    return cv2.warpPerspective(img, Hm, (S, S), flags=cv2.INTER_LINEAR, borderMode=cv2.BORDER_REPLICATE)


def rectify_legacy(name: str, out_size=None):
    img = load_legacy(name)
    corners, info = detect_window(img)
    return rectify_quad(img, corners, out_size), corners, info


# ------------------------------------------------------------------------- duplicate check
def _dhash_variants(gray64: np.ndarray):
    """dHash (16x16 -> 256 bits) of the 8 dihedral variants of a square grey thumbnail."""
    out = []
    for k in range(4):
        r = np.rot90(gray64, k)
        for flip in (False, True):
            g = r[:, ::-1] if flip else r
            s = cv2.resize(g, (17, 16), interpolation=cv2.INTER_AREA)
            out.append((s[:, 1:] > s[:, :-1]).ravel())
    return out


def duplicate_check(rectified150: dict, legacy_crops: dict):
    """Min Hamming distance (over 8 dihedral variants) of dHash + mean-colour distance.

    Returns a list of (legacy_name, best_150_id, hamming/256, colour_dist) sorted by hamming.
    Unrelated textures give ~0.5; near duplicates < 0.2.
    """
    def thumb(bgr):
        return cv2.resize(cv2.cvtColor(bgr, cv2.COLOR_BGR2GRAY), (64, 64), interpolation=cv2.INTER_AREA)

    def lab_mean(bgr):
        return cv2.cvtColor(cv2.resize(bgr, (32, 32), interpolation=cv2.INTER_AREA),
                            cv2.COLOR_BGR2LAB).reshape(-1, 3).mean(0)

    h150 = {i: _dhash_variants(thumb(c))[0] for i, c in rectified150.items()}
    c150 = {i: lab_mean(c) for i, c in rectified150.items()}
    res = []
    for n, crop in legacy_crops.items():
        hv = _dhash_variants(thumb(crop))
        cm = lab_mean(crop)
        best = None
        for i, h in h150.items():
            d = min(np.mean(h != v) for v in hv)
            cd = float(np.linalg.norm(cm - c150[i]))
            if best is None or d < best[2]:
                best = (n, i, float(d), cd)
        res.append(best)
    return sorted(res, key=lambda t: t[2])


if __name__ == "__main__":
    for r in load_gt():
        print(r)
    print("2019 baseline (weft):", load_baseline2019())
