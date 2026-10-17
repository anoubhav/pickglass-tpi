"""
run_pipeline.py -- FAITHFUL baseline harness for the legacy IIT-M "Textile Quality Analysis"
warp/weft counting pipeline, evaluated on the 20 ground-truth images in
"$PICKGLASS_LEGACY_ROOT/FR-final attempt/images/".

Nothing here re-tunes or improves the algorithm. Source files are only READ.

Pipeline reproduced (see fidelity notes in the docstrings below):

  [segmentation]  The final counting code (TQA/doc3.py) was run on PRE-CROPPED fabric
                  images (TQA/images f1..f23 = TQA/croppedimages.zip, no resize, doc3.py
                  L6-9 + comment L18 "Assume the center of image always contains fabric").
                  The FR-final ground-truth images a1..r10 are likewise already cropped to
                  the 1-inch pick-glass window, so the PRIMARY variant ("precropped") feeds
                  them straight into doc3.  A SECONDARY variant ("seg_final") first runs the
                  extracted notebook segmentation (seg_final.segment) and feeds its crop to doc3.

  [dot map]       TQA/doc3.py L9-L88, transcribed line by line (same NumPy-1.x uint8
                  semantics, same OpenCV calls / parameters; only GUI calls dropped).

  [count]         TQA/count.py L5-L12 (row y=0, `img[y,x]-img[y,x-1]==255` on uint8 of the
                  JPEG-reloaded 80x80 dot map) -> "warp".  count.py has no weft pass; the
                  report says counting is done "row and column wise", so weft = the same
                  test mirrored to column x=0.

  [per inch]      count.py counts over the 80-px width of the dot map (doc3 k=40 crop).
                  GT is threads in the 1-inch window; the project itself treated a whole
                  FR image as the 1-inch window (test.py / manual.py count across the full
                  image and the xlsx compares that directly with GT).  So
                  warp_per_inch = warp80 * W_img/80,  weft_per_inch = weft80 * H_img/80.

Usage (venv):
  venv/bin/python run_pipeline.py                      # both variants, seed 0, + seed sweep
  venv/bin/python run_pipeline.py --variants precropped --seed-sweep 0
"""
import argparse
import contextlib
import csv
import datetime as _dt
import io
import json
import os
import re
import sys
import time
import warnings

import cv2
import numpy as np

warnings.filterwarnings("ignore", category=DeprecationWarning)
warnings.filterwarnings("ignore", category=RuntimeWarning)  # uint8 scalar wrap in count.py is intended
from scipy.ndimage import interpolation as inter  # noqa: E402  (doc3.py L3, still present in SciPy 1.x)

import openpyxl  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import seg_final  # noqa: E402

FA = (os.environ.get("PICKGLASS_LEGACY_ROOT") or os.path.join(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."), "data", "legacy2019"))
DEF_IMAGES = os.path.join(FA, "FR-final attempt", "images")
DEF_GT = os.path.join(FA, "Warp-Weft count.xlsx")
DEF_GT2 = os.path.join(FA, "FR-final attempt", "Warp-Weft count.xlsx")
TQA = os.path.join(FA, "TQA")
ORDER = [f"a{i}" for i in range(1, 11)] + [f"r{i}" for i in range(1, 11)]


# =====================================================================================
# 1. TQA/doc3.py  (dot-map generation) -- transcribed verbatim, line numbers in comments
# =====================================================================================
def doc3_dotmap(img):
    """Body of the doc3.py for-loop for ONE image (L9-L85). `img` is BGR uint8 as from
    cv2.imread (doc3 L8, no resize). Returns dict with `final` (80x80 uint8, L85), the
    full-frame binary map (L84), best rotation angle, k-means centres and printed line.
    Deviations: imshow/waitKey dropped (headless); the L84 C argument is wrapped in float()
    because OpenCV 4.11 does not accept a 1-element ndarray there (same value)."""
    imgray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)                                      # L9
    thresh = cv2.adaptiveThreshold(imgray, 1, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                                   cv2.THRESH_BINARY, 21, 0)                             # L12
    thresh = cv2.blur(thresh, (3, 3))                                                    # L13
    k = 40                                                                               # L19
    h, w = thresh.shape                                                                  # L20
    filter = thresh[int(h / 2 - k):int(h / 2 + k), int(w / 2 - k):int(w / 2 + k)]      # L22
    filter = 2 * filter - 1          # L24  (uint8: 0 -> 255 wrap, 1 -> 1)
    filter = filter * 255 / (4 * k * k)   # L25  (uint8 *255 wraps: 255->1, 1->255; then /6400)
    filteri = filter * -1                                                                # L26
    thresh = thresh * 2 - np.ones(np.shape(thresh))                                     # L28 (float64 +-1)
    filtered_img = cv2.filter2D(thresh, cv2.CV_8UC3, filter)                             # L31
    filtered_imgi = cv2.filter2D(thresh, cv2.CV_8UC3, filteri)                           # L32
    filtered_img = cv2.subtract(filtered_img, filtered_imgi)                             # L33

    def find_score(img, angle):                                                          # L36
        data = inter.rotate(img, angle, reshape=False, order=0)                          # L37
        hist = np.sum(data, axis=1)                                                      # L38
        score = np.sum((hist[1:] - hist[:-1]) ** 2)                                      # L39
        return hist, score                                                               # L40

    inc = 1                                                                              # L42
    max_angle = 10                                                                       # L43
    angles = range(-max_angle, max_angle, inc)                                           # L44
    scores = []                                                                          # L45
    for angle in angles:                                                                 # L47
        hist, score = find_score(filtered_img, angle)                                    # L48
        scores.append(score)                                                             # L49
    best_score = max(scores)                                                             # L51
    best_angle = angles[scores.index(best_score)]                                        # L52
    filtered_img = inter.rotate(filtered_img, best_angle, reshape=False, order=0)        # L54
    kernel = np.ones((3, 3), np.uint8)                                                   # L57
    filtered_img2 = filtered_img                                                         # L59
    filtered_img2 = filtered_img2 - cv2.blur(filtered_img2, (51, 51)) + 125              # L61 (uint8 wrap)
    filtered_img = cv2.dilate(filtered_img, kernel, iterations=1)                        # L63
    Z = filtered_img2.reshape((-1, 1))                                                   # L65
    Z = np.float32(Z)                                                                    # L67
    criteria = (cv2.TERM_CRITERIA_EPS + cv2.TERM_CRITERIA_MAX_ITER, 10, 1.0)             # L69
    K = 3                                                                                # L70
    ret, label, center = cv2.kmeans(Z, K, None, criteria, 10, cv2.KMEANS_RANDOM_CENTERS)  # L71
    for i in range(3):                                                                   # L73
        if (np.max(center) == center[i]):                                                # L74
            a = i                                                                        # L75
        elif (np.min(center) == center[i]):                                              # L76
            c = i                                                                        # L77
        else:                                                                            # L78
            b = i                                                                        # L79
    printed = f"{center[a]} {center[b]} {center[c]}"                                     # L80 (print)
    centers_abc = (float(center[a][0]), float(center[b][0]), float(center[c][0]))
    if (center[b] < 128):                                                                # L81
        b = a                                                                            # L82
    C = float(-np.abs(center[c] - center[b]) * 1.3 / 2)
    filtered_img = cv2.adaptiveThreshold(filtered_img, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                                         cv2.THRESH_BINARY, 51, C)                       # L84
    final = filtered_img[int(h / 2 - k):int(h / 2 + k), int(w / 2 - k):int(w / 2 + k)]  # L85
    return {"final": final, "fullmap": filtered_img, "best_angle": int(best_angle),
            "centers_abc": centers_abc, "C": C, "printed": printed}


# =====================================================================================
# 2. TQA/count.py  (counting) -- verbatim, plus the column mirror for weft
# =====================================================================================
def count_py_warp(img):
    """TQA/count.py L6-L12 (identical to commented doc3.py L91-97). img = uint8 grayscale
    as returned by cv2.imread('final.jpg', 0). uint8 scalar arithmetic: the ==255 test is
    true for a 0->255 rise AND for any -1 step (e.g. 255->254 JPEG noise); x=0 compares
    with x=-1 (wrap-around)."""
    h, w = img.shape[:2]                                                                 # L6
    warp1 = 0                                                                            # L7
    for y in range(0, 1):                                                                # L9
        for x in range(0, w):                                                            # L10
            if img[y, x] - img[y, x - 1] == 255:                                         # L11
                warp1 += 1                                                               # L12
    return warp1


def count_py_weft_mirror(img):
    """NOT in the original: count.py's test mirrored to column x=0 scanning down y
    (the report: 'counting the periodic patterns row and column wise')."""
    h, w = img.shape[:2]
    weft1 = 0
    for x in range(0, 1):
        for y in range(0, h):
            if img[y, x] - img[y - 1, x] == 255:
                weft1 += 1
    return weft1


def diag_rising_edges(img):
    """Diagnostic only (not used for metrics): median over ALL rows / ALL columns of clean
    0->255 rises (binarised at 127, no wrap-around) in the same 80x80 dot map."""
    bb = (img > 127).astype(np.int8)
    rows = (np.diff(bb, axis=1) == 1).sum(axis=1)
    cols = (np.diff(bb, axis=0) == 1).sum(axis=0)
    return float(np.median(rows)), float(np.median(cols))


# =====================================================================================
# 3. Ground truth
# =====================================================================================
def _parse_count(v):
    """-> (value_for_error, low, high, raw_text). Text formulas 'a + b*2 = N' -> N (value
    after the last '='); ranges 'lo-hi' -> midpoint; numbers as-is; None -> None."""
    if v is None:
        return None, None, None, ""
    if isinstance(v, (int, float)):
        return float(v), float(v), float(v), str(v)
    s = str(v).strip()
    if "=" in s:
        n = float(s.split("=")[-1].strip())
        return n, n, n, s
    m = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)\s*", s)
    if m:
        lo, hi = float(m.group(1)), float(m.group(2))
        return (lo + hi) / 2, lo, hi, s
    if s == "" or s == "*":
        return None, None, None, s
    return float(s), float(s), float(s), s


def _parse_weave(v):
    """Excel turned '1/1' and '2/1' into dates (number_format 'm/d'): month/day -> weave."""
    if v is None:
        return ""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return f"{v.month}/{v.day}"
    return str(v).strip()


def load_gt(path, path2=None):
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr) if h is not None}
    gt = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = row[col["Img file name"]]
        if name is None or not re.fullmatch(r"[ar]\d+\*?", str(name).strip()):
            continue
        name = str(name).strip()
        key = name.rstrip("*")
        weft = _parse_count(row[col["weft count"]])
        warp = _parse_count(row[col["warp count"]])
        gt[key] = {
            "unclear": name.endswith("*"),
            "weft": weft[0], "weft_lo": weft[1], "weft_hi": weft[2], "weft_raw": weft[3],
            "warp": warp[0], "warp_lo": warp[1], "warp_hi": warp[2], "warp_raw": warp[3],
            "weave": _parse_weave(row[col["interlacing pattern"]]),
            "thread_count_raw": "" if row[col["thread count"]] is None else str(row[col["thread count"]]),
            "usable": weft[0] is not None and warp[0] is not None,
        }
    xcheck = []
    if path2 and os.path.exists(path2):
        ws2 = openpyxl.load_workbook(path2, data_only=False).worksheets[0]
        hdr2 = [c.value for c in ws2[1]]
        col2 = {h: i for i, h in enumerate(hdr2) if h is not None}
        for row in ws2.iter_rows(min_row=2, values_only=True):
            name = row[col2["Img file name"]]
            if name is None or not re.fullmatch(r"[ar]\d+\*?", str(name).strip()):
                continue
            key = str(name).strip().rstrip("*")
            g = gt.get(key)
            w2 = _parse_count(row[col2["weft count"]])[0]
            p2 = _parse_count(row[col2["warp count"]])[0]
            wv2 = _parse_weave(row[col2["interlacing pattern"]])
            same = g is not None and g["weft"] == w2 and g["warp"] == p2 and g["weave"] == wv2
            if not same:
                xcheck.append(key)
            if g is not None:
                rgb = row[col2["predicted weft by RGB irregularity"]]
                g["xlsx_rgb_weft_pred"] = None if rgb is None else float(rgb)
    return gt, xcheck


# =====================================================================================
# 4. Per-image pipeline
# =====================================================================================
def _fname_for(key, images_dir):
    for n in os.listdir(images_dir):
        if os.path.splitext(n)[0].strip() == key and n.lower().endswith((".jpg", ".jpeg", ".png")):
            return n
    return None


def run_one(path, fname, variant, seed, out_dir, save=True):
    """Returns dict of results for one image/variant."""
    original = cv2.imread(path)                                   # doc3 L8 (no resize)
    H, W = original.shape[:2]
    rec = {"image": fname, "variant": variant, "seed": seed, "img_h": H, "img_w": W}
    if variant == "precropped":
        doc3_in = original
        rec.update(seg_status="skipped (input already cropped to window)", seg_box="")
    elif variant == "seg_final":
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            crop, box, fb = seg_final.segment(path, cv_seed=seed)
        rec.update(seg_status="fallback" if fb else "valid", seg_box=json.dumps(np.asarray(box).tolist()),
                   seg_stdout=buf.getvalue().strip().replace("\n", " | "))
        if crop is None:
            rec.update(error="segmentation returned degenerate box")
            return rec
        doc3_in = crop
    else:
        raise ValueError(variant)
    rec["doc3_in_h"], rec["doc3_in_w"] = doc3_in.shape[:2]
    cv2.setRNGSeed(int(seed))          # NOT in the original (unseeded cv2.kmeans); for repeatability
    try:
        r = doc3_dotmap(doc3_in)
    except NameError as e:             # e.g. tied k-means centres leave `b` undefined (doc3 L73-79)
        rec.update(error=f"doc3 NameError: {e}")
        return rec
    sub = os.path.join(out_dir, variant if save else f"_sweep/{variant}_seed{seed}")
    os.makedirs(os.path.join(sub, "dotmaps"), exist_ok=True)
    # doc3 L88 writes ./outputs/{fname}.jpg (JPEG); count.py L5 reads it back as grayscale
    jpg = os.path.join(sub, "dotmaps", f"{fname}.jpg")
    cv2.imwrite(jpg, r["final"])
    final_reloaded = cv2.imread(jpg, 0)
    warp80 = count_py_warp(final_reloaded)
    weft80 = count_py_weft_mirror(final_reloaded)
    # doc3.py L91-97 (commented-out) applies the same test to the in-memory `final` array
    # (no JPEG round trip) -> recorded as a secondary original-code variant
    warp80_mem = count_py_warp(r["final"])
    weft80_mem = count_py_weft_mirror(r["final"])
    d_row, d_col = diag_rising_edges(final_reloaded)
    rec.update(warp_count_80px_inmem=warp80_mem, weft_count_80px_inmem=weft80_mem,
               pred_warp_per_inch_inmem=round(warp80_mem * W / 80, 2),
               pred_weft_per_inch_inmem=round(weft80_mem * H / 80, 2))
    rec.update(best_angle=r["best_angle"], kmeans_centers_max_mid_min=json.dumps([round(x, 2) for x in r["centers_abc"]]),
               at2_C=round(r["C"], 3), dotmap_white_frac=round(float((final_reloaded > 127).mean()), 4),
               warp_count_80px=warp80, weft_count_80px=weft80,
               px_per_inch_x=W, px_per_inch_y=H,
               crop_extent_in_x=round(80 / W, 4), crop_extent_in_y=round(80 / H, 4),
               pred_warp_per_inch=round(warp80 * W / 80, 2), pred_weft_per_inch=round(weft80 * H / 80, 2),
               diag_median_row_rises_80px=d_row, diag_median_col_rises_80px=d_col,
               dotmap_jpg=jpg)
    if save:
        os.makedirs(os.path.join(sub, "fullmaps"), exist_ok=True)
        os.makedirs(os.path.join(sub, "panels"), exist_ok=True)
        stem = fname.rsplit(".", 1)[0].strip()
        cv2.imwrite(os.path.join(sub, "fullmaps", f"{stem}_fullmap.png"), r["fullmap"])
        _panel(doc3_in, r["fullmap"], final_reloaded, fname, rec, os.path.join(sub, "panels", f"{stem}_panel.jpg"))
    return rec


def _job(a):
    """Pool worker: (path, fname, variant, seed, out_dir, save) -> rec. Each job seeds the
    cv2 RNG itself, so results do not depend on process or order."""
    return run_one(*a)


def _run_jobs(jobs, workers):
    if workers <= 1:
        return [_job(j) for j in jobs]
    import multiprocessing as mp
    with mp.get_context("spawn").Pool(workers) as pool:
        return pool.map(_job, jobs, chunksize=1)


def _panel(inp, fullmap, final, fname, rec, out):
    S = 480
    a = cv2.resize(inp, (S, S), interpolation=cv2.INTER_AREA)
    b = cv2.cvtColor(cv2.resize(fullmap, (S, S), interpolation=cv2.INTER_AREA), cv2.COLOR_GRAY2BGR)
    for im in (a, b):  # mark the 80x80 analysed centre region
        hh, ww = (inp.shape[0], inp.shape[1])
        x0, y0 = int((ww / 2 - 40) * S / ww), int((hh / 2 - 40) * S / hh)
        x1, y1 = int((ww / 2 + 40) * S / ww), int((hh / 2 + 40) * S / hh)
        cv2.rectangle(im, (x0, y0), (x1, y1), (0, 0, 255), 2)
    c = cv2.cvtColor(cv2.resize(final, (S, S), interpolation=cv2.INTER_NEAREST), cv2.COLOR_GRAY2BGR)
    cv2.line(c, (0, 2), (S - 1, 2), (0, 0, 255), 3)      # row y=0 (warp, count.py)
    cv2.line(c, (2, 0), (2, S - 1), (255, 0, 0), 3)      # col x=0 (weft, mirror)
    canvas = np.concatenate([a, np.full((S, 6, 3), 255, np.uint8), b, np.full((S, 6, 3), 255, np.uint8), c], axis=1)
    bar = np.full((60, canvas.shape[1], 3), 255, np.uint8)
    cv2.putText(bar, f"{fname} [{rec['variant']}] warp80={rec['warp_count_80px']} weft80={rec['weft_count_80px']} "
                     f"-> warp/in={rec['pred_warp_per_inch']} weft/in={rec['pred_weft_per_inch']}  angle={rec['best_angle']}",
                (8, 38), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (0, 0, 0), 2)
    cv2.imwrite(out, np.concatenate([bar, canvas], axis=0))


# =====================================================================================
# 5. Metrics
# =====================================================================================
def metrics(rows, gt, pred_key, gt_key, keys):
    errs, pes, within = [], [], 0
    for k in keys:
        p = rows[k].get(pred_key)
        g = gt[k][gt_key]
        if p is None or g is None:
            continue
        e = abs(p - g)
        errs.append(e)
        pes.append(100 * e / g)
        within += int(e <= 0.10 * g)
    n = len(errs)
    return {"n": n, "MAE": round(float(np.mean(errs)), 2) if n else None,
            "MAPE_pct": round(float(np.mean(pes)), 1) if n else None,
            "within_10pct": within,
            "median_abs_err": round(float(np.median(errs)), 2) if n else None}


def selftest_tqa(out_dir, seed):
    """Fidelity check of the transcription: re-run doc3 on TQA/images and compare with the
    saved TQA/outputs/{f}.jpg.jpg; run count.py on TQA/final.jpg."""
    res = []
    names = sorted([n for n in os.listdir(os.path.join(TQA, "images")) if n.endswith(".jpg")],
                   key=lambda s: int(s[1:-4]))
    sd = os.path.join(out_dir, "selftest_tqa_dotmaps")
    os.makedirs(sd, exist_ok=True)
    for n in names:
        cv2.setRNGSeed(seed)
        r = doc3_dotmap(cv2.imread(os.path.join(TQA, "images", n)))
        saved_path = os.path.join(TQA, "outputs", n + ".jpg")
        saved = cv2.imread(saved_path, 0)
        ours_path = os.path.join(sd, n + ".jpg")
        cv2.imwrite(ours_path, r["final"])                      # doc3 L88 (JPEG)
        ours = cv2.imread(ours_path, 0)                          # count.py L5
        agree = float(((r["final"] > 127) == (saved > 127)).mean())
        res.append({"image": n, "pix_agree_vs_saved_output": round(agree, 4),
                    "jpeg_bytes_identical_to_saved": open(ours_path, "rb").read() == open(saved_path, "rb").read(),
                    "count_py_on_saved_jpeg": count_py_warp(saved),
                    "count_py_on_our_jpeg": count_py_warp(ours),
                    "count_py_on_our_array_no_jpeg": count_py_warp(r["final"]),
                    "clean_rises_row0_saved": int((np.diff((saved[0] > 127).astype(np.int8)) == 1).sum())})
    with open(os.path.join(out_dir, "selftest_tqa_doc3.csv"), "w", newline="") as f:
        wr = csv.DictWriter(f, fieldnames=list(res[0].keys()))
        wr.writeheader()
        wr.writerows(res)
    final_count = count_py_warp(cv2.imread(os.path.join(TQA, "final.jpg"), 0))
    ag = [r["pix_agree_vs_saved_output"] for r in res]
    return {"count_py_on_TQA_final.jpg": final_count,
            "doc3_pix_agree_median": round(float(np.median(ag)), 4),
            "doc3_pix_agree_min": round(float(np.min(ag)), 4),
            "n_ge_0.98": int(sum(a >= 0.98 for a in ag)), "n": len(ag),
            "n_jpeg_bytes_identical": int(sum(r["jpeg_bytes_identical_to_saved"] for r in res)),
            "n_count_py_equal_on_jpeg": int(sum(r["count_py_on_saved_jpeg"] == r["count_py_on_our_jpeg"] for r in res))}


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--images", default=DEF_IMAGES)
    ap.add_argument("--gt", default=DEF_GT)
    ap.add_argument("--gt2", default=DEF_GT2)
    ap.add_argument("--out", default=os.path.join(HERE, "run1"))
    ap.add_argument("--variants", default="precropped,seg_final")
    ap.add_argument("--seed", type=int, default=0)
    ap.add_argument("--seed-sweep", type=int, default=5, help="extra cv2 RNG seeds for k-means variability (precropped)")
    ap.add_argument("--no-selftest", action="store_true")
    ap.add_argument("--workers", type=int, default=6, help="parallel processes (results are per-job seeded)")
    args = ap.parse_args(argv)
    os.makedirs(args.out, exist_ok=True)
    t0 = time.time()

    gt, xcheck = load_gt(args.gt, args.gt2)
    usable = [k for k in ORDER if k in gt and gt[k]["usable"]]
    clear = [k for k in usable if not gt[k]["unclear"]]
    print(f"GT: {len(usable)} usable ({len(clear)} clear); cross-check mismatches vs FR xlsx: {xcheck}")

    summary = {"gt_file": args.gt, "gt_crosscheck_file": args.gt2, "gt_crosscheck_mismatches": xcheck,
               "n_usable": len(usable), "usable": usable, "n_clear": len(clear), "clear": clear,
               "seed": args.seed, "variants": {}}
    if not args.no_selftest:
        summary["selftest_tqa"] = selftest_tqa(args.out, args.seed)
        print("selftest:", summary["selftest_tqa"])

    variants = [v.strip() for v in args.variants.split(",") if v.strip()]
    fnames = {key: _fname_for(key, args.images) for key in ORDER}
    for variant in variants:
        rows = {}
        jobs = [(os.path.join(args.images, fnames[k]), fnames[k], variant, args.seed, args.out, True)
                for k in ORDER if fnames[k] is not None]
        done = {r["image"]: r for r in _run_jobs(jobs, args.workers)}
        for key in ORDER:
            fname = fnames[key]
            if fname is None:
                rows[key] = {"image": key, "variant": variant, "error": "missing file"}
                continue
            rec = done[fname]
            g = gt.get(key, {})
            rec.update(key=key, gt_weave=g.get("weave", ""), gt_unclear=g.get("unclear", ""),
                       gt_warp_raw=g.get("warp_raw", ""), gt_weft_raw=g.get("weft_raw", ""),
                       gt_warp=g.get("warp"), gt_weft=g.get("weft"),
                       gt_warp_range=(f"{g['warp_lo']:g}-{g['warp_hi']:g}" if g.get("warp") is not None and g['warp_lo'] != g['warp_hi'] else ""),
                       gt_usable=g.get("usable", False),
                       xlsx_rgb_weft_pred=g.get("xlsx_rgb_weft_pred"))
            for d in ("warp", "weft"):
                p, gv = rec.get(f"pred_{d}_per_inch"), g.get(d)
                if p is not None and gv is not None:
                    rec[f"abs_err_{d}"] = round(abs(p - gv), 2)
                    rec[f"pct_err_{d}"] = round(100 * abs(p - gv) / gv, 1)
                    rec[f"within10_{d}"] = abs(p - gv) <= 0.1 * gv
            rows[key] = rec
            print(f"[{variant}] {fname:9s} warp80={rec.get('warp_count_80px')} weft80={rec.get('weft_count_80px')} "
                  f"-> warp/in={rec.get('pred_warp_per_inch')} weft/in={rec.get('pred_weft_per_inch')} "
                  f"GT warp={g.get('warp')} weft={g.get('weft')} seg={rec.get('seg_status')} angle={rec.get('best_angle')}",
                  flush=True)
        fields = []
        for r in rows.values():
            for f in r:
                if f not in fields:
                    fields.append(f)
        csv_name = "results.csv" if variant == "precropped" else f"results_{variant}.csv"
        with open(os.path.join(args.out, csv_name), "w", newline="") as f:
            wr = csv.DictWriter(f, fieldnames=fields)
            wr.writeheader()
            for k in ORDER:
                wr.writerow(rows[k])
        m = {}
        for subset_name, subset in (("usable", usable), ("clear_only", clear)):
            m[subset_name] = {
                "warp": metrics(rows, gt, "pred_warp_per_inch", "warp", subset),
                "weft": metrics(rows, gt, "pred_weft_per_inch", "weft", subset)}
            m[subset_name + "__inmem_doc3_L91_97"] = {
                "warp": metrics(rows, gt, "pred_warp_per_inch_inmem", "warp", subset),
                "weft": metrics(rows, gt, "pred_weft_per_inch_inmem", "weft", subset)}
        m["seg_fallbacks"] = sum(1 for r in rows.values() if r.get("seg_status") == "fallback")
        m["errors"] = {k: r["error"] for k, r in rows.items() if "error" in r}
        summary["variants"][variant] = m
        print(variant, json.dumps(m))

    # reference: the project's own 'predicted weft by RGB irregularity' column (test.py)
    ref_rows = {k: {"pred": gt[k].get("xlsx_rgb_weft_pred")} for k in usable}
    summary["reference_xlsx_rgb_weft_pred"] = {
        "usable": metrics(ref_rows, gt, "pred", "weft", usable),
        "clear_only": metrics(ref_rows, gt, "pred", "weft", clear)}

    # k-means seed sweep (cv2.kmeans KMEANS_RANDOM_CENTERS is unseeded in doc3.py)
    if args.seed_sweep > 0 and "precropped" in variants:
        sweep = []
        per_seed = {}
        seeds = list(range(args.seed + 1, args.seed + 1 + args.seed_sweep))
        jobs = [(os.path.join(args.images, fnames[k]), fnames[k], "precropped", s, args.out, False)
                for s in seeds for k in ORDER]
        allrec = _run_jobs(jobs, args.workers)
        for s in seeds:
            rows = {}
            for key in ORDER:
                rec = next(r for r in allrec if r["seed"] == s and r["image"] == fnames[key])
                rows[key] = rec
                sweep.append({"seed": s, "key": key, "warp80": rec.get("warp_count_80px"),
                              "weft80": rec.get("weft_count_80px"),
                              "pred_warp_per_inch": rec.get("pred_warp_per_inch"),
                              "pred_weft_per_inch": rec.get("pred_weft_per_inch"),
                              "warp80_inmem": rec.get("warp_count_80px_inmem"),
                              "weft80_inmem": rec.get("weft_count_80px_inmem"),
                              "best_angle": rec.get("best_angle"), "error": rec.get("error", "")})
            per_seed[s] = {"warp": metrics(rows, gt, "pred_warp_per_inch", "warp", usable),
                           "weft": metrics(rows, gt, "pred_weft_per_inch", "weft", usable)}
        with open(os.path.join(args.out, "seed_sweep_precropped.csv"), "w", newline="") as f:
            wr = csv.DictWriter(f, fieldnames=list(sweep[0].keys()))
            wr.writeheader()
            wr.writerows(sweep)
        summary["seed_sweep_precropped_usable"] = per_seed
    summary["runtime_s"] = round(time.time() - t0, 1)
    summary["versions"] = {"python": sys.version.split()[0], "numpy": np.__version__, "cv2": cv2.__version__,
                           "scipy": __import__("scipy").__version__}
    with open(os.path.join(args.out, "metrics.json"), "w") as f:
        json.dump(summary, f, indent=1)
    print("wrote", os.path.join(args.out, "metrics.json"), "runtime", summary["runtime_s"], "s")
    return 0


if __name__ == "__main__":
    sys.exit(main())
